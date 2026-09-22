"""Inventory valuation and annual historical-return scenarios."""


def value_inventory(mean_stock: float, unit_cost_ars: float, ars_per_usd: float) -> float:
    """Value a stock quantity at a fixed acquisition cost and ARS/USD rate."""
    if ars_per_usd <= 0:
        raise ValueError("The ARS per USD exchange rate must be positive")
    return mean_stock * unit_cost_ars / ars_per_usd


def opportunity_cost(stock_value_usd: float, annual_return: float) -> float:
    """Annual USD scenario, with a decimal total return (not a percent)."""
    return stock_value_usd * annual_return


def read_inventory(path):
    """Reconstruct TP2 states and retain the owner's literal cost inputs."""
    import re
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook["Hoja1"]
    # Materialize once: repeated random access to read-only XLSX is expensive.
    cells = dict(enumerate(sheet.iter_rows(min_row=15, max_row=73, max_col=5), start=15))
    def cell(row, column):
        return cells[row][column - 1].value

    pallet_bags = int(re.search(r"=\s*(\d+)", cell(59, 1)).group(1))
    components = [{"source_cell": f"E{row}", "label": cell(row, 1), "amount_ars": cell(row, 5)}
                  for row in range(59, 64)]
    pallet_cost = sum(item["amount_ars"] for item in components)
    if pallet_cost != cell(65, 5):
        raise ValueError("Owner cost subtotal does not reconcile")
    opening = cell(15, 5)
    balance = opening
    purchases = sales = 0
    audit = [{"source_row": 15, "purchase": 0, "sale": 0,
              "original_balance": opening, "corrected_balance": opening, "adjustment": 0}]
    for row in range(16, 52):
        purchase = cell(row, 3)
        sale = cell(row, 4)
        if purchase is None and sale is None:
            raise ValueError(f"No movement in ledger row {row}")
        purchase = 0 if purchase is None else purchase
        sale = 0 if sale is None else sale
        balance += purchase - sale
        purchases += purchase
        sales += sale
        original = cell(row, 5)
        audit.append({"source_row": row, "purchase": purchase, "sale": sale,
                      "original_balance": original, "corrected_balance": balance,
                      "adjustment": None if original is None else balance - original})
    workbook.close()
    stock_sum = sum(item["corrected_balance"] for item in audit)
    return {"pallet_bags": pallet_bags, "pallet_cost_ars": pallet_cost,
            "unit_cost_ars": pallet_cost / pallet_bags, "cost_components": components,
            "factory_unit_cost_ars": components[0]["amount_ars"] / pallet_bags,
            "opening_stock": opening, "closing_stock": balance,
            "purchases": purchases, "sales": sales, "stock_audit": audit,
            "stock_sum": stock_sum, "stock_count": len(audit), "mean_stock": stock_sum / len(audit)}


def validate_funds(funds):
    """Only compare same-date, same-period USD NAV total returns."""
    bases = {(f["currency"], f["as_of_date"], f["period"], f["return_type"]) for f in funds}
    if len(funds) != 2 or len(bases) != 1:
        raise ValueError("Fund observations are not comparable")
    currency, _, period, return_type = next(iter(bases))
    if (currency, period, return_type) != ("USD", "1y", "navSourced"):
        raise ValueError("A comparable annual USD NAV total return is required")


def analyze(root):
    """Reproduce all report values offline from immutable local evidence."""
    import hashlib
    import json
    from datetime import date

    raw = root / "data/raw"
    reference = root / "data/reference"
    workbook = raw / "inventory.xlsx"
    if hashlib.sha256(workbook.read_bytes()).hexdigest() != "240f403b1c5755b6fdd5082ad4c3b720358eafa7f6475508405d9f9773277f85":
        raise ValueError("Original inventory workbook changed")
    result = read_inventory(workbook)
    market = json.loads((reference / "market-assumptions.json").read_text())
    funds = market["funds"]
    validate_funds(funds)
    exchange = market["exchange_rate"]
    for fund in funds:
        source = json.loads((reference / fund["evidence_file"]).read_text())
        points = source["componentsByNameMap"]["performance"]["containersByNameMap"]["returns"]["subContainersByNameMap"]["average"]["dataPointsByNameMap"]
        index = points["returnTypes"]["value"].index("navSourced")
        if (points["ticker"]["value"] != fund["ticker"]
            or str(points["asOfDate"]["value"]) != fund["as_of_date"].replace("-", "")
            or str(points["returnTypes"]["asOfDate"][index]) != fund["as_of_date"].replace("-", "")
            or source["currencyCode"] != fund["currency"]
            or points["oneYearAnnualized"]["value"][index] != fund["annual_total_return_percent"]
            or float(fund["annual_total_return_percent"]) / 100 != fund["annual_return"]):
            raise ValueError("Pinned market assumption differs from issuer evidence")
        if date.fromisoformat(fund["as_of_date"]) > date.fromisoformat(exchange["date"]):
            raise ValueError("Fund return observation follows valuation date")
    value = value_inventory(result["mean_stock"], result["unit_cost_ars"], exchange["ars_per_usd"])
    result.update({"exchange_rate": exchange, "funds": funds,
                   "unit_cost_usd": result["unit_cost_ars"] / exchange["ars_per_usd"],
                   "stock_value_ars": result["mean_stock"] * result["unit_cost_ars"],
                   "stock_value_usd": value, "report_date": market["report_date"]})
    for fund in funds:
        fund["annual_cost_usd"] = opportunity_cost(value, fund["annual_return"])
        fund["unit_annual_cost_usd"] = opportunity_cost(result["unit_cost_usd"], fund["annual_return"])
    result["annual_cost_difference_usd"] = funds[0]["annual_cost_usd"] - funds[1]["annual_cost_usd"]
    return result
